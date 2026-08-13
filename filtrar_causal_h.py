"""
filtrar_causal_h.py

Filtra rechazos_2026.xlsx para identificar SOLO los rechazos fundados
(total o parcialmente) en la causal del Articulo 20 letra h) de la LPI
(anterioridades / marcas previas), que es la causal que el buscador de
anterioridades esta destinado a predecir.

Motivo: rechazos_2026.xlsx contiene TODOS los tipos de rechazo (letra e,
f, k, h, etc.). Medir el recall del buscador contra el total de rechazos
sin filtrar por causal produce un falso negativo: el buscador no tiene
por que detectar una marca rechazada por ser generica (letra e) o por
orden publico (letra k), asi que ese caso cuenta como "fallo" cuando en
realidad esta fuera de alcance del buscador.

Regla de conteo (definida con Cami, 17-jul-2026):
- Si un rechazo cita la causal h) junto con otra causal (ej. e y h juntas),
  se cuenta igual como "causal h", porque el buscador SI deberia haber
  detectado esa anterioridad, sin importar que ademas haya otras razones
  de rechazo.
- Deduplicacion por File_Nbr: una solicitud puede tener varias filas (una
  por cada clase NCL cubierta). Si CUALQUIER fila de esa solicitud cita
  causal h), la solicitud completa se cuenta como rechazo por causal h.

Uso:
    python filtrar_causal_h.py --input data/rechazos_2026.xlsx --output data/rechazos_causal_h.xlsx

Dependencias: pandas, openpyxl (ya instaladas en el entorno buscador)
"""

import argparse
import re
import sys
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# Normalizacion de texto (el campo Notes1 viene con encoding roto: sin
# tildes, con literales "nbs;" insertados donde deberia haber espacios no
# separables, y basura HTML residual en algunos casos)
# ---------------------------------------------------------------------------

def normalizar_texto(texto):
    if pd.isna(texto):
        return ""
    t = str(texto).lower()
    t = t.replace("nbs;", " ")
    t = unicodedata.normalize("NFKD", t)
    t = "".join(c for c in t if not unicodedata.combining(c))
    t = re.sub(r"\s+", " ", t)
    return t


# ---------------------------------------------------------------------------
# Deteccion de causales citadas en el texto de la observacion / resolucion
# ---------------------------------------------------------------------------

# Letra individual seguida de parentesis de cierre, ej. "h)". Se usa dentro
# de listas del tipo "letra e), f) y k)" donde solo la primera letra viene
# precedida de la palabra "letra".
LETRA_SUELTA = re.compile(r"\b([a-k])\s*\)")

# Un "bloque de causales" es "letra" seguido de una o mas letras entre
# parentesis, separadas por coma o "y": cubre tanto el caso simple
# ("letra h)") como listas ("letra e), f) y k)" o "letra e) y f)").
BLOQUE_CAUSALES = re.compile(
    r"letra\s*(?:[a-k]\s*\)\s*(?:,\s*|y\s*)?)+", re.IGNORECASE
)

CONTEXTO_ART20 = re.compile(r"art[i]?culos?\.?\s*(?:19\s*(?:y|,)\s*)?20\b")


def detectar_causales(texto_normalizado):
    """
    Devuelve el set de letras de causal citadas en el texto.

    Busca bloques que empiezan con la palabra "letra" y luego contienen una
    o mas letras entre parentesis (cubriendo listas como "letra e), f) y k)"
    o "letra e) y f)"), y exige que el bloque este en el contexto de una
    referencia al articulo 20 dentro de una ventana razonable hacia atras.
    """
    causales = set()
    for bloque in BLOQUE_CAUSALES.finditer(texto_normalizado):
        inicio_ventana = max(0, bloque.start() - 60)
        ventana = texto_normalizado[inicio_ventana:bloque.start()]
        if CONTEXTO_ART20.search(ventana):
            letras = LETRA_SUELTA.findall(bloque.group())
            causales.update(letras)
    return causales


REGISTRO_PATRON = re.compile(r"registro\s*n?\.?\s*[oº]?\s*(\d{5,9})")


def extraer_registros_citados(texto_normalizado):
    """Extrae los numeros de Registro citados como marca(s) previa(s) opuesta(s)."""
    return sorted(set(REGISTRO_PATRON.findall(texto_normalizado)))


# ---------------------------------------------------------------------------
# Procesamiento principal
# ---------------------------------------------------------------------------

def procesar(input_path: Path, output_path: Path):
    print(f"Leyendo {input_path} ...")
    df = pd.read_excel(input_path)

    columnas_esperadas = ["File_Nbr", "Nice_Class_Description", "Mark_Name",
                           "Status_Name", "Notes1", "Notification_Date"]
    faltantes = [c for c in columnas_esperadas if c not in df.columns]
    if faltantes:
        print(f"ADVERTENCIA: faltan columnas esperadas: {faltantes}")
        print(f"Columnas disponibles: {list(df.columns)}")

    df["_texto_norm"] = df["Notes1"].apply(normalizar_texto)
    df["_causales"] = df["_texto_norm"].apply(detectar_causales)
    df["causal_h"] = df["_causales"].apply(lambda s: "h" in s)
    df["causales_detectadas"] = df["_causales"].apply(lambda s: ", ".join(sorted(s)) if s else "")
    df["registros_citados"] = df["_texto_norm"].apply(
        lambda t: ", ".join(extraer_registros_citados(t))
    )
    df["tiene_multiples_causales"] = df["_causales"].apply(lambda s: len(s) > 1)

    # -----------------------------------------------------------------
    # Agregacion por File_Nbr (una solicitud puede tener varias filas,
    # una por clase NCL). Una solicitud cuenta como "causal h" si
    # CUALQUIERA de sus filas cita causal h).
    # -----------------------------------------------------------------
    agg = df.groupby("File_Nbr").agg(
        Mark_Name=("Mark_Name", "first"),
        Status_Name=("Status_Name", "first"),
        Notification_Date=("Notification_Date", "first"),
        clases_cubiertas=("Nice_Class_Description", lambda s: " | ".join(
            x[:80] for x in s.dropna().astype(str).unique()
        )),
        causal_h=("causal_h", "any"),
        causales_detectadas=("causales_detectadas", lambda s: ", ".join(
            sorted(set(c.strip() for v in s if v for c in v.split(",") if c.strip()))
        )),
        registros_citados=("registros_citados", lambda s: ", ".join(
            sorted(set(r.strip() for v in s if v for r in v.split(",") if r.strip()))
        )),
        filas_asociadas=("File_Nbr", "count"),
    ).reset_index()

    agg["tiene_multiples_causales"] = agg["causales_detectadas"].apply(
        lambda s: len([c for c in s.split(",") if c.strip()]) > 1
    )

    total_marcas = len(agg)
    total_causal_h = int(agg["causal_h"].sum())
    total_solo_h = int((agg["causal_h"] & ~agg["tiene_multiples_causales"]).sum())
    total_h_mas_otra = int((agg["causal_h"] & agg["tiene_multiples_causales"]).sum())
    total_sin_causal_detectada = int((agg["causales_detectadas"] == "").sum())

    print()
    print("=" * 60)
    print("RESUMEN")
    print("=" * 60)
    print(f"Total solicitudes rechazadas (File_Nbr unicos): {total_marcas}")
    print(f"Rechazadas citando causal h) (sola o con otras):  {total_causal_h}  ({100*total_causal_h/total_marcas:.1f}%)")
    print(f"  - Solo causal h):                               {total_solo_h}")
    print(f"  - Causal h) + otra causal:                      {total_h_mas_otra}")
    print(f"Rechazadas SIN causal h) (otras causales):        {total_marcas - total_causal_h}  ({100*(total_marcas-total_causal_h)/total_marcas:.1f}%)")
    if total_sin_causal_detectada:
        print()
        print(f"ATENCION: {total_sin_causal_detectada} solicitudes no tuvieron ninguna causal "
              f"detectada por el regex (Notes1 vacio, formato distinto, o texto no cubierto "
              f"por el patron). Revisar manualmente en la pestana 'Sin_Causal_Detectada' del Excel.")
    print("=" * 60)

    # -----------------------------------------------------------------
    # Exportar a Excel
    # -----------------------------------------------------------------
    solo_h = agg[agg["causal_h"]].copy().sort_values("File_Nbr")
    sin_h = agg[~agg["causal_h"]].copy().sort_values("File_Nbr")
    sin_causal_detectada = agg[agg["causales_detectadas"] == ""].copy()

    resumen_data = {
        "Metrica": [
            "Total solicitudes rechazadas (base)",
            "Rechazadas por causal h) (sola o combinada)",
            "  Solo causal h)",
            "  Causal h) + otra causal",
            "Rechazadas SIN causal h)",
            "Sin causal detectada por regex (revisar manualmente)",
            "% rechazos que SI corresponden a causal h) sobre el total",
        ],
        "Valor": [
            total_marcas,
            total_causal_h,
            total_solo_h,
            total_h_mas_otra,
            total_marcas - total_causal_h,
            total_sin_causal_detectada,
            f"{100*total_causal_h/total_marcas:.1f}%" if total_marcas else "N/A",
        ],
    }
    df_resumen = pd.DataFrame(resumen_data)

    columnas_salida = [
        "File_Nbr", "Mark_Name", "Status_Name", "Notification_Date",
        "causal_h", "tiene_multiples_causales", "causales_detectadas",
        "registros_citados", "clases_cubiertas", "filas_asociadas",
    ]

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_resumen.to_excel(writer, sheet_name="Estadisticas", index=False)
        solo_h[columnas_salida].to_excel(writer, sheet_name="Causal_H", index=False)
        sin_h[columnas_salida].to_excel(writer, sheet_name="Otras_Causales", index=False)
        if len(sin_causal_detectada):
            sin_causal_detectada[columnas_salida].to_excel(
                writer, sheet_name="Sin_Causal_Detectada", index=False
            )

    _formatear_workbook(output_path)

    print(f"\nArchivo generado: {output_path}")
    print(f"  - Hoja 'Estadisticas': resumen general")
    print(f"  - Hoja 'Causal_H': {len(solo_h)} solicitudes para usar en la validacion del buscador")
    print(f"  - Hoja 'Otras_Causales': {len(sin_h)} solicitudes fuera de alcance del buscador")

    return agg


def _formatear_workbook(path: Path):
    """Aplica formato basico: encabezados en negrita, fuente Arial, columnas autoajustadas."""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    body_font = Font(name="Arial", size=10)

    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font

        for col_cells in ws.columns:
            max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
            col_letter = get_column_letter(col_cells[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 60)

        ws.freeze_panes = "A2"

    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description="Filtra rechazos por causal Art. 20 letra h)")
    parser.add_argument("--input", type=str, default="data/rechazos_2026.xlsx",
                         help="Ruta al Excel de rechazos (default: data/rechazos_2026.xlsx)")
    parser.add_argument("--output", type=str, default="data/rechazos_causal_h.xlsx",
                         help="Ruta de salida (default: data/rechazos_causal_h.xlsx)")
    args = parser.parse_args()

    input_path = Path(args.input)
    output_path = Path(args.output)

    if not input_path.exists():
        print(f"ERROR: no se encontro el archivo de entrada: {input_path}")
        sys.exit(1)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    procesar(input_path, output_path)


if __name__ == "__main__":
    main()
